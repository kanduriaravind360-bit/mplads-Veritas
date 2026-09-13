"""Check that the numbers quoted in the docs and the deck match their sources.

    python scripts/check_numbers.py

Each claim is recomputed from ``models/metrics.json`` (always) and the app
database (when ``data/app/sentinel.db`` exists), formatted the way the documents
write it, and looked for in every document that quotes it. A missing claim means
a document has drifted from the system. Exit code 1 if any claim is missing.
"""

from __future__ import annotations

import json
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
README = "README.md"
QA = "docs/JUDGE_QA.md"
DEMO = "docs/DEMO_SCRIPT.md"
DECK = "presentation/MPLADS_Sentinel_SIH26102.pptx"


@dataclass(frozen=True)
class Claim:
    label: str
    text: str
    documents: tuple[str, ...]


def pct(x: float) -> str:
    return f"{x * 100:.0f}%"


def indian(n: float) -> str:
    s = f"{int(round(n))}"
    if len(s) <= 3:
        return s
    head, tail = s[:-3], s[-3:]
    groups = []
    while len(head) > 2:
        groups.insert(0, head[-2:])
        head = head[:-2]
    if head:
        groups.insert(0, head)
    return ",".join([*groups, tail])


def metric_claims(m: dict) -> list[Claim]:
    inj = m["injection_test"]["per_stream_recall"]
    holdout = m["holdout_evaluation"]
    delay = holdout["delay_model"]["holdout"]
    cost = m["expected_cost_model"]
    return [
        Claim(
            "split detector fired",
            pct(inj["split_group"]["detector_fired"]),
            (README, QA, DEMO, DECK),
        ),
        Claim(
            "cost top-5% recall",
            pct(inj["inflated_cost"]["recall_at_top_5pct"]),
            (README, QA, DECK),
        ),
        Claim(
            "duplicate top-5% recall",
            pct(inj["duplicate"]["recall_at_top_5pct"]),
            (README, QA, DECK),
        ),
        Claim("holdout delay ROC-AUC", f"{delay['roc_auc']:.3f}", (README, QA, DEMO, DECK)),
        Claim("holdout delay PR-AUC", f"{delay['pr_auc']:.3f}", (README, QA, DECK)),
        Claim(
            "held-out constituencies", str(holdout["n_holdout_constituencies"]), (QA, DEMO, DECK)
        ),
        Claim("held-out works", indian(holdout["n_holdout_works"]), (QA, DECK)),
        Claim("proxy PR-AUC", f"{m['supervised_model']['pr_auc_oof']:.3f}", (README, QA)),
        Claim("expected-cost typical error", f"{cost['mae_as_cost_ratio']:.1f}", (QA, DECK)),
        Claim("expected-cost R²", f"{cost['r2']:.2f}", (QA, DECK)),
        Claim("severe floor works lifted", indian(m["severe_floor"]["works_lifted"]), (QA,)),
        Claim("duplicate pairs", indian(m["duplicates"]["pairs"]), (DECK,)),
        Claim("split groups", indian(m["split_works"]["groups"]), (QA, DECK)),
    ]


def database_claims() -> list[Claim]:
    db_path = ROOT / "data" / "app" / "sentinel.db"
    if not db_path.exists():
        return []
    sys.path.insert(0, str(ROOT))
    from sqlalchemy import func, select

    from backend.app import db as dbm
    from backend.app import models, queries
    from backend.app.scoping import Scope

    dbm.configure(dbm.make_engine(f"sqlite:///{db_path.as_posix()}"))
    with dbm.session_factory()() as session:
        t = queries.totals(session, Scope(role="MINISTRY"))
        alerts = session.execute(
            select(func.count()).select_from(models.Alert).where(models.Alert.is_active.is_(True))
        ).scalar_one()
        district = (
            session.execute(
                select(models.User).where(models.User.role == "DISTRICT").order_by(models.User.id)
            )
            .scalars()
            .first()
        )
        district_works = (
            queries.totals(session, Scope.for_user(district))["works"] if district else 0
        )
    return [
        Claim("works", indian(t["works"]), (README, QA, DEMO, DECK)),
        Claim("sanctioned crore", f"{t['sanctioned'] / 1e7:,.0f} crore", (README, DEMO, DECK)),
        Claim("high or critical works", indian(t["high_or_critical"]), (README, QA, DEMO, DECK)),
        Claim(
            "money at risk crore", f"{t['money_at_risk'] / 1e7:,.0f} crore", (README, DEMO, DECK)
        ),
        Claim("open alerts", indian(alerts), (QA, DECK)),
        Claim("district officer works", indian(district_works), (QA, DECK)),
    ]


def document_text(relative: str) -> str:
    path = ROOT / relative
    if path.suffix == ".pptx":
        from pptx import Presentation

        parts = []
        for slide in Presentation(str(path)).slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    parts.extend(c.text for row in shape.table.rows for c in row.cells)
        return "\n".join(parts)
    return path.read_text(encoding="utf-8")


def check(claims: list[Claim]) -> list[str]:
    texts: dict[str, str] = {}
    missing = []
    for claim in claims:
        for doc in claim.documents:
            if doc not in texts:
                texts[doc] = document_text(doc)
            if claim.text not in texts[doc]:
                missing.append(f"{doc}: '{claim.text}' ({claim.label}) not found")
    return missing


def main() -> int:
    metrics = json.loads((ROOT / "models" / "metrics.json").read_text(encoding="utf-8"))
    claims = metric_claims(metrics) + database_claims()
    missing = check(claims)
    checks = sum(len(c.documents) for c in claims)
    print(f"{len(claims)} claims, {checks} document checks")
    for line in missing:
        print("  MISSING " + line)
    print("all quoted numbers match their sources" if not missing else f"{len(missing)} mismatches")
    return 1 if missing else 0


if __name__ == "__main__":
    sys.exit(main())
